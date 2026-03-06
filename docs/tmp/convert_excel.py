#!/usr/bin/env python3
"""Convert Excel files (.xlsx/.xls) in docs/ to Markdown + CSV hybrid format."""

import csv
import io
import os
import re
import shutil
import sys
from datetime import date
from pathlib import Path

import openpyxl
import xlrd

DOCS_DIR = Path(__file__).resolve().parent.parent  # docs/
ARCHIVE_DIR = DOCS_DIR / "_archive"
TODAY = date.today().isoformat()

# Files to embed as markdown tables (small/borderline)
EMBED_FILES = {
    "Advancement/arrowoflightadvancementreport-0.xlsx",
    "Advancement/bearadvancementreport0-0.xlsx",
    "Advancement/wolfadvancementreport0.xlsx",
    "Advancement/tigeradvancementreport0-0.xlsx",
    "Advancement/lionsadvancementreport.xlsx",
    "Advancement/webelosadvancementreport0.xlsx",
    "Events & Activities/Pinewood Derby/2016/2016Standings.xlsx",
    "Events & Activities/Pinewood Derby/2017/2017standings.xlsx",
    "Membership/2025 Pack 232 AOL Crossover.xlsx",
    "Membership/Family Talent Survey/2025/Cub_Scout_Parent_Talent_Survey.xlsx",
    "Fundraising/Popcorn/2013/PackPopcorn232ShowandSellWalmart2013-0.xlsx",
    "Fundraising/Popcorn/2013/PackPopcorn232ShowandSellWalmart2013.xlsx",
}

# Files that are reference-only (complex dashboards)
REFERENCE_FILES = {
    "Fundraising/Popcorn/2024/2024 Popcorn Dashboard.xlsx",
    "Fundraising/Popcorn/2025/2025 Popcorn Dashboard.xlsx",
}

# Everything else gets CSV treatment

# Category mapping based on directory
CATEGORY_MAP = {
    "Advancement": "advancement",
    "Events & Activities": "events",
    "Fundraising": "fundraising",
    "Leadership": "leadership",
    "Membership": "membership",
    "Training": "training",
    "Treasurer": "financial",
}


def get_category(rel_path: str) -> str:
    top_dir = rel_path.split("/")[0]
    return CATEGORY_MAP.get(top_dir, "general")


def get_tags(rel_path: str, category: str) -> list:
    tags = [category]
    parts = rel_path.split("/")
    if "Pinewood Derby" in rel_path:
        tags.append("pinewood-derby")
    if "Popcorn" in rel_path:
        tags.append("popcorn")
    if "Leader Awards" in rel_path:
        tags.append("leader-awards")
    if "Leader" in rel_path and "leader-awards" not in tags:
        tags.append("leadership")
    if "Talent Survey" in rel_path:
        tags.append("talent-survey")
    return list(dict.fromkeys(tags))  # deduplicate preserving order


def make_title(filename: str) -> str:
    name = Path(filename).stem
    # Clean up common patterns
    name = re.sub(r"[-_]", " ", name)
    name = re.sub(r"(\d)([A-Z])", r"\1 \2", name)
    name = re.sub(r"([a-z])([A-Z])", r"\1 \2", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name


def make_description(title: str, rel_path: str) -> str:
    category = get_category(rel_path)
    return f"Pack 232 {category} spreadsheet: {title}"


def archive_path(rel_path: str) -> Path:
    return ARCHIVE_DIR / rel_path


def relative_archive_link(md_path: Path, arch_path: Path) -> str:
    """Compute relative path from md file to archive file."""
    return os.path.relpath(arch_path, md_path.parent)


def sanitize_cell(val) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    # Escape pipe characters in markdown tables
    s = s.replace("|", "\\|")
    # Replace newlines with spaces
    s = s.replace("\n", " ").replace("\r", "")
    return s


def read_xlsx_sheets(filepath: Path) -> dict:
    """Read all sheets from an xlsx file. Returns {sheet_name: [[row], ...]}"""
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    sheets = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([sanitize_cell(c) for c in row])
        sheets[name] = rows
    wb.close()
    return sheets


def read_xls_sheets(filepath: Path) -> dict:
    """Read all sheets from a legacy .xls file."""
    wb = xlrd.open_workbook(str(filepath))
    sheets = {}
    for name in wb.sheet_names():
        ws = wb.sheet_by_name(name)
        rows = []
        for r in range(ws.nrows):
            rows.append([sanitize_cell(ws.cell_value(r, c)) for c in range(ws.ncols)])
        sheets[name] = rows
    return sheets


def read_sheets(filepath: Path) -> dict:
    if filepath.suffix.lower() == ".xls":
        return read_xls_sheets(filepath)
    return read_xlsx_sheets(filepath)


def strip_empty_rows(rows: list) -> list:
    """Remove trailing fully-empty rows."""
    while rows and all(c == "" for c in rows[-1]):
        rows.pop()
    return rows


def rows_to_markdown_table(rows: list) -> str:
    if not rows:
        return "_Empty sheet_"
    rows = strip_empty_rows(rows)
    if not rows:
        return "_Empty sheet_"

    # First row is header
    header = rows[0]
    ncols = len(header)

    # Pad rows to same width
    for i, row in enumerate(rows):
        if len(row) < ncols:
            rows[i] = row + [""] * (ncols - len(row))
        elif len(row) > ncols:
            ncols = len(row)
            header = header + [""] * (len(row) - len(header))

    # If header is all empty, use generic column names
    if all(h == "" for h in header):
        header = [f"Column {i+1}" for i in range(ncols)]
        data_rows = rows  # include first row as data
    else:
        data_rows = rows[1:]

    lines = []
    lines.append("| " + " | ".join(header) + " |")
    lines.append("| " + " | ".join(["---"] * ncols) + " |")
    for row in data_rows:
        padded = row + [""] * (ncols - len(row))
        lines.append("| " + " | ".join(padded[:ncols]) + " |")

    return "\n".join(lines)


def rows_to_csv_string(rows: list) -> str:
    rows = strip_empty_rows(rows)
    buf = io.StringIO()
    writer = csv.writer(buf)
    for row in rows:
        writer.writerow(row)
    return buf.getvalue()


def frontmatter(title: str, description: str, source_file: str,
                archive_link: str, fmt: str, category: str, tags: list,
                extra: dict = None) -> str:
    tag_str = ", ".join(f'"{t}"' for t in tags)
    lines = [
        "---",
        f'title: "{title}"',
        f'description: "{description}"',
        f'source_file: "{source_file}"',
        f'source_archive: "{archive_link}"',
        f'converted_date: "{TODAY}"',
        f'original_format: "{fmt}"',
        f'author: "Unknown"',
        f'tags: [{tag_str}]',
        f'category: "{category}"',
    ]
    if extra:
        for k, v in extra.items():
            if isinstance(v, str):
                lines.append(f'{k}: "{v}"')
            else:
                lines.append(f'{k}: {v}')
    lines.append("draft: false")
    lines.append("---")
    return "\n".join(lines)


def process_embed(filepath: Path, rel_path: str):
    """Embed small spreadsheet as markdown table."""
    sheets = read_sheets(filepath)
    title = make_title(filepath.name)
    desc = make_description(title, rel_path)
    category = get_category(rel_path)
    tags = get_tags(rel_path, category)
    fmt = filepath.suffix.lstrip(".").lower()

    md_path = filepath.with_suffix(".md")
    arch = archive_path(rel_path)
    arch_link = relative_archive_link(md_path, arch)

    body_parts = []
    fm = frontmatter(title, desc, filepath.name, arch_link, fmt, category, tags)
    body_parts.append(fm)
    body_parts.append("")

    # If multiple sheets, add headers for each
    sheet_names = list(sheets.keys())
    for sname in sheet_names:
        rows = sheets[sname]
        if len(sheet_names) > 1:
            body_parts.append(f"## {sname}")
            body_parts.append("")
        body_parts.append(rows_to_markdown_table(rows))
        body_parts.append("")

    body_parts.append("---")
    body_parts.append("")
    body_parts.append(f"> **Original spreadsheet**: [{filepath.name}]({arch_link})")
    body_parts.append("")

    md_path.write_text("\n".join(body_parts), encoding="utf-8")
    print(f"  EMBED: {md_path.name}")

    # Archive original
    do_archive(filepath, arch)


def process_csv(filepath: Path, rel_path: str):
    """Convert to CSV companion files + summary .md."""
    sheets = read_sheets(filepath)
    title = make_title(filepath.name)
    desc = make_description(title, rel_path)
    category = get_category(rel_path)
    tags = get_tags(rel_path, category)
    fmt = filepath.suffix.lstrip(".").lower()

    md_path = filepath.with_suffix(".md")
    arch = archive_path(rel_path)
    arch_link = relative_archive_link(md_path, arch)

    fm = frontmatter(title, desc, filepath.name, arch_link, fmt, category, tags)
    body_parts = [fm, ""]
    body_parts.append(f"# {title}")
    body_parts.append("")

    sheet_names = list(sheets.keys())
    if len(sheet_names) > 1:
        body_parts.append("## Sheets")
        body_parts.append("")

    basename = filepath.stem
    for sname in sheet_names:
        rows = sheets[sname]
        rows = strip_empty_rows(rows)
        if not rows:
            continue

        # Determine CSV filename
        if len(sheet_names) == 1:
            csv_name = f"{basename}.csv"
        else:
            csv_name = f"{basename} - {sname}.csv"

        csv_path = filepath.parent / csv_name
        csv_content = rows_to_csv_string(rows)
        csv_path.write_text(csv_content, encoding="utf-8")
        print(f"  CSV:   {csv_name}")

        # Count stats
        nrows = len(rows) - 1  # exclude header
        ncols = max(len(r) for r in rows) if rows else 0
        header_preview = ", ".join(rows[0][:5]) if rows else ""
        if len(rows[0]) > 5:
            header_preview += ", ..."

        body_parts.append(f"- **[{sname}](./{csv_name})** - {nrows} rows, {ncols} columns. Columns: {header_preview}")

    body_parts.append("")
    body_parts.append("---")
    body_parts.append("")
    body_parts.append(f"> **Original spreadsheet**: [{filepath.name}]({arch_link})")
    body_parts.append("")

    md_path.write_text("\n".join(body_parts), encoding="utf-8")
    print(f"  MD:    {md_path.name}")

    do_archive(filepath, arch)


def process_reference(filepath: Path, rel_path: str):
    """Create reference-only .md for complex dashboards."""
    title = make_title(filepath.name)
    desc = make_description(title, rel_path)
    category = get_category(rel_path)
    tags = get_tags(rel_path, category)
    fmt = filepath.suffix.lstrip(".").lower()

    md_path = filepath.with_suffix(".md")
    arch = archive_path(rel_path)
    arch_link = relative_archive_link(md_path, arch)

    # Try to get sheet names for reference
    try:
        sheets = read_sheets(filepath)
        sheet_names = list(sheets.keys())
        sheet_count = len(sheet_names)
    except Exception:
        sheet_names = []
        sheet_count = "unknown"

    fm = frontmatter(title, desc, filepath.name, arch_link, fmt, category, tags,
                     extra={"type": "reference"})
    body_parts = [fm, ""]
    body_parts.append(f"# {title}")
    body_parts.append("")
    body_parts.append(f"This is a complex multi-sheet Excel dashboard ({sheet_count} sheets) that cannot be fully represented in text format.")
    body_parts.append("Please refer to the original spreadsheet for full details.")
    body_parts.append("")

    if sheet_names:
        body_parts.append("## Sheet Names")
        body_parts.append("")
        for sn in sheet_names:
            body_parts.append(f"- {sn}")
        body_parts.append("")

    body_parts.append("---")
    body_parts.append("")
    body_parts.append(f"> **Original spreadsheet**: [{filepath.name}]({arch_link})")
    body_parts.append("")

    md_path.write_text("\n".join(body_parts), encoding="utf-8")
    print(f"  REF:   {md_path.name}")

    do_archive(filepath, arch)


def do_archive(src: Path, dst: Path):
    """Archive file using copy2 + unlink (WSL2 safe)."""
    dst.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(str(src), str(dst))
    src.unlink()
    print(f"  ARCH:  -> {dst.relative_to(DOCS_DIR)}")


def main():
    # Find all Excel files
    excel_files = []
    for ext in ("*.xlsx", "*.xls"):
        for f in DOCS_DIR.rglob(ext):
            if "_archive" not in f.parts:
                excel_files.append(f)
    excel_files.sort()

    print(f"Found {len(excel_files)} Excel files to convert\n")

    embed_count = 0
    csv_count = 0
    ref_count = 0
    skip_count = 0

    for filepath in excel_files:
        rel_path = str(filepath.relative_to(DOCS_DIR))
        print(f"Processing: {rel_path}")

        # Check if .md already exists
        md_path = filepath.with_suffix(".md")
        if md_path.exists():
            print(f"  SKIP:  .md already exists")
            skip_count += 1
            continue

        # Determine treatment
        if rel_path in REFERENCE_FILES:
            process_reference(filepath, rel_path)
            ref_count += 1
        elif rel_path in EMBED_FILES:
            process_embed(filepath, rel_path)
            embed_count += 1
        else:
            process_csv(filepath, rel_path)
            csv_count += 1

    print(f"\n--- Summary ---")
    print(f"Embedded:    {embed_count}")
    print(f"CSV:         {csv_count}")
    print(f"Reference:   {ref_count}")
    print(f"Skipped:     {skip_count}")
    print(f"Total:       {embed_count + csv_count + ref_count + skip_count}")


if __name__ == "__main__":
    main()
